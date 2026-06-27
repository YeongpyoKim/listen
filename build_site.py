# -*- coding: utf-8 -*-
"""
build_site.py — 경청 프로젝트 '동네 한 바퀴' 사이트 데이터 빌더
"""
import json
import re
import shutil
from pathlib import Path

import openpyxl

from site_content import CONTENT

ROOT = Path(__file__).resolve().parent
XLSX = ROOT / "상점 리스트.xlsx"
SITE = ROOT / "site"
ASSETS = SITE / "assets"
DATA = SITE / "data"
NAVER_MENUS = DATA / "naver_menus.json"

IMG_EXT = {".jpg", ".jpeg", ".png", ".webp", ".gif"}

COLS = [
    "name", "type", "manager", "signature", "price", "hours",
    "seats", "baby_chair", "parking", "history", "tip", "charm",
]


def natural_key(path: Path):
    m = re.search(r"(\d+)", path.stem)
    return (0, int(m.group(1))) if m else (1, path.stem)


def slugify(index: int) -> str:
    return f"s{index:02d}"


def read_excel():
    wb = openpyxl.load_workbook(XLSX)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=4, values_only=False):
        name_cell = row[1]
        if not name_cell.value:
            continue
        rec = {}
        for i, key in enumerate(COLS):
            val = row[1 + i].value
            rec[key] = ("" if val is None else str(val).strip())
        rec["naver_url"] = name_cell.hyperlink.target if name_cell.hyperlink else ""
        rows.append(rec)
    return rows


def collect_images(store_name: str):
    folder = ROOT / store_name
    if not folder.is_dir():
        return [], []
    imgs = [p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in IMG_EXT]
    scene, menu = [], []
    for p in imgs:
        if "메뉴" in p.name:
            menu.append(p)
        else:
            scene.append(p)
    scene.sort(key=natural_key)
    menu.sort(key=natural_key)
    return scene, menu


def load_naver_menus():
    if not NAVER_MENUS.exists():
        return {}
    try:
        return json.loads(NAVER_MENUS.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return {}


def parse_menu_items(signature: str, price: str):
    if not signature:
        return []
    items = []
    for part in re.split(r"[,，\n/]+", signature):
        name = part.strip()
        if not name or len(name) > 60:
            continue
        item = {"name": name}
        if price and len(price) < 40:
            item["price"] = price
        items.append(item)
    return items[:12]


def pick_menu_items(name, signature, price, naver_menus):
    naver = naver_menus.get(name, {})
    items = naver.get("items") or []
    if items:
        return items
    return parse_menu_items(signature, price)


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    DATA.mkdir(parents=True, exist_ok=True)

    excel_rows = read_excel()
    naver_menus = load_naver_menus()
    stores = []

    for idx, rec in enumerate(excel_rows, start=1):
        name = rec["name"]
        sid = slugify(idx)
        content = CONTENT.get(name, {})

        scene_imgs, menu_imgs = collect_images(name)
        dest_dir = ASSETS / sid
        gallery = []
        menu_images = []
        main_image = ""

        if dest_dir.exists():
            shutil.rmtree(dest_dir)
        if scene_imgs or menu_imgs:
            dest_dir.mkdir(parents=True, exist_ok=True)

            for i, src in enumerate(scene_imgs):
                ext = src.suffix.lower()
                fname = f"{i:02d}{ext}"
                shutil.copy2(src, dest_dir / fname)
                gallery.append(f"assets/{sid}/{fname}")

            for src in menu_imgs:
                stem = re.sub(r"[^\w가-힣]", "", src.stem) or "menu"
                fname = f"{stem}{src.suffix.lower()}"
                dest = dest_dir / fname
                n = 1
                while dest.exists():
                    fname = f"{stem}{n}{src.suffix.lower()}"
                    dest = dest_dir / fname
                    n += 1
                shutil.copy2(src, dest_dir / fname)
                menu_images.append(f"assets/{sid}/{fname}")

            main_image = gallery[0] if gallery else ""

        menu_items = pick_menu_items(name, rec["signature"], rec["price"], naver_menus)

        store = {
            "id": sid,
            "name": name,
            "type": rec["type"],
            "emoji": content.get("emoji", "🍽️"),
            "accent": content.get("accent", "#8a6d3b"),
            "tagline": content.get("tagline", ""),
            "story": content.get("story", ""),
            "sunday": content.get("sunday", "check"),
            "sunday_note": content.get("sunday_note", ""),
            "address": content.get("address", ""),
            "phone": content.get("phone", ""),
            "signature": rec["signature"],
            "price": rec["price"],
            "hours": rec["hours"],
            "seats": rec["seats"],
            "baby_chair": rec["baby_chair"],
            "parking": rec["parking"],
            "history": rec["history"],
            "tip": rec["tip"],
            "charm": rec["charm"],
            "naver_url": rec["naver_url"],
            "main_image": main_image,
            "gallery": gallery,
            "menu_images": menu_images,
            "menu_items": menu_items,
            "placeholder": main_image == "",
        }
        if content.get("sunday_rule"):
            store["sunday_rule"] = content["sunday_rule"]
        stores.append(store)

    out = {
        "title": "동네 한 바퀴, 은혜 한 바구니",
        "subtitle": "경청 프로젝트 2기 · 작은 실천",
        "count": len(stores),
        "stores": stores,
    }
    (DATA / "stores.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    placeholders = [s["name"] for s in stores if s["placeholder"]]
    print(f"[OK] {len(stores)}개 상점 → {DATA / 'stores.json'}")
    print(f"[i] 사진 폴더 없는(감성 카드) 상점 {len(placeholders)}개: {', '.join(placeholders)}")


if __name__ == "__main__":
    main()
